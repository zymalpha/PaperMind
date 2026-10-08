from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Protocol

import fitz
from docx import Document

from .models import DocumentPage


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class DocumentLoader(Protocol):
    def load(self, path: Path, document_id: str) -> list[DocumentPage]: ...


def calculate_document_id(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


class PDFLoader:
    def load(self, path: Path, document_id: str) -> list[DocumentPage]:
        pages: list[DocumentPage] = []
        with fitz.open(path) as pdf:
            if pdf.needs_pass:
                raise ValueError(f"文档已加密，无法解析: {path.name}")
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text", sort=True).strip()
                if text:
                    pages.append(_page(text, path, page_number, document_id))
        return pages


class DocxLoader:
    def load(self, path: Path, document_id: str) -> list[DocumentPage]:
        document = Document(path)
        blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
        text = "\n\n".join(blocks)
        return [_page(text, path, 1, document_id)] if text else []


class TextLoader:
    def load(self, path: Path, document_id: str) -> list[DocumentPage]:
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                text = path.read_text(encoding=encoding).strip()
                return [_page(text, path, 1, document_id)] if text else []
            except UnicodeDecodeError:
                continue
        raise ValueError(f"无法识别文本编码: {path.name}")


class LoaderRegistry:
    def __init__(self) -> None:
        self._loaders: dict[str, DocumentLoader] = {
            ".pdf": PDFLoader(),
            ".docx": DocxLoader(),
            ".txt": TextLoader(),
            ".md": TextLoader(),
        }

    def load(self, path: str | Path) -> tuple[str, list[DocumentPage]]:
        source = Path(path).resolve()
        if not source.is_file():
            raise FileNotFoundError(f"文件不存在: {source}")
        suffix = source.suffix.lower()
        if suffix not in self._loaders:
            supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
            raise ValueError(f"不支持 {suffix or '无扩展名'} 文件，支持: {supported}")
        document_id = calculate_document_id(source)
        pages = self._loaders[suffix].load(source, document_id)
        if not pages:
            raise ValueError(f"文档中没有可索引的文本: {source.name}")
        return document_id, pages


def _page(text: str, path: Path, page: int, document_id: str) -> DocumentPage:
    return DocumentPage(
        text=text,
        source_path=str(path),
        source_name=path.name,
        page=page,
        document_id=document_id,
    )

