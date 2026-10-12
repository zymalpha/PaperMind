"""Render delivery Markdown files into editable Word documents using python-docx."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DELIVERY = ROOT / "完整提交材料"
DOCUMENTS = [
    "系统评测报告.md",
    "Bad Case分析报告.md",
    "用户使用手册.md",
    "技术设计文档.md",
    "提交材料目录说明.md",
]


def set_run_font(run, size: float = 10.5, bold: bool | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    properties.append(shade)


def add_inline(paragraph, text: str, size: float = 10.5) -> None:
    token_pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
    for part in token_pattern.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size, True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size - 0.5)
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.color.rgb = RGBColor(55, 65, 81)
        else:
            set_run_font(paragraph.add_run(part), size)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    header = rows[0]
    data_rows = [row for row in rows[1:] if not all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in row)]
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, content in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = ""
        set_cell_shading(cell, "DCE6F1")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(paragraph, content.strip(), 9.2)
        for run in paragraph.runs:
            run.bold = True
    for row in data_rows:
        if len(row) < len(header):
            row = row + [""] * (len(header) - len(row))
        cells = table.add_row().cells
        for index, content in enumerate(row[: len(header)]):
            cells[index].text = ""
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            add_inline(cells[index].paragraphs[0], content.strip(), 8.8)
    document.add_paragraph()


def add_footer(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("PaperMind · 南京农业大学生产实习    |    "), 8.5)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def render_markdown(source: Path, target: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    styles = document.styles
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Heading 1"].font.color.rgb = RGBColor(31, 54, 83)
    styles["Heading 2"].font.color.rgb = RGBColor(38, 78, 119)
    styles["Heading 3"].font.color.rgb = RGBColor(64, 85, 106)
    add_footer(document)

    index = 0
    code_lines: list[str] = []
    in_code = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.22)
                paragraph.paragraph_format.space_after = Pt(7)
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                run.font.size = Pt(8.6)
                in_code, code_lines = False, []
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            table_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(document, table_rows)
            continue
        image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            image_path = (source.parent / image_match.group(2)).resolve()
            if image_path.is_file():
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(image_path), width=Inches(6.1))
                caption = document.add_paragraph(image_match.group(1))
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in caption.runs:
                    set_run_font(run, 9)
            else:
                paragraph = document.add_paragraph()
                add_inline(paragraph, f"[图表未找到：{image_path}]")
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            paragraph = document.add_paragraph(style="Title" if level == 1 else f"Heading {level - 1}")
            add_inline(paragraph, heading.group(2), 20 if level == 1 else 14 if level == 2 else 12)
            if level == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            index += 1
            continue
        if stripped.startswith("> "):
            paragraph = document.add_paragraph(style="Quote")
            add_inline(paragraph, stripped[2:])
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, bullet.group(1))
            index += 1
            continue
        numbered = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(1))
            index += 1
            continue
        paragraph = document.add_paragraph()
        add_inline(paragraph, stripped)
        index += 1

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)


def main() -> int:
    missing = [name for name in DOCUMENTS if not (DELIVERY / name).is_file()]
    if missing:
        raise FileNotFoundError(f"尚未生成交付 Markdown：{', '.join(missing)}")
    for name in DOCUMENTS:
        source = DELIVERY / name
        target = source.with_suffix(".docx")
        render_markdown(source, target)
        print(f"created {target.relative_to(ROOT)} ({target.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
